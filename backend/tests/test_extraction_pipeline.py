"""Real parser/render fixtures with a deterministic provider double, never live-model accuracy claims."""
from __future__ import annotations

import codecs
import hashlib
import io
import json
import shutil
import zipfile
from copy import deepcopy
from types import SimpleNamespace

import fitz
import pytest
from docx import Document
from PIL import Image

from northstar_api.services import extraction_regions as regions
from northstar_api.services.extraction_checkpoint import decode_page_checkpoint, encode_page_checkpoint, read_zip
from northstar_api.services.extraction_pipeline import CheckpointWriteError, extract_document
from northstar_api.services.extraction_word import inspect_word, render_word, xml

MODEL = "meta/llama-3.2-11b-vision-instruct"


class TestSettings(SimpleNamespace):
    __test__ = False

    def model_dump(self) -> dict:
        return vars(self).copy()


def settings(**overrides: object) -> TestSettings:
    values = dict(extraction_vision_model=MODEL, extraction_vision_provider="nvidia", extraction_vision_timeout_seconds=10,
                  extraction_vision_max_tokens=512, extraction_render_dpi=72, extraction_max_image_pixels=1_000_000,
                  extraction_max_calls_per_page=16, extraction_region_retries=1, extraction_max_evidence_bytes=10_000_000,
                  extraction_word_renderer="soffice", extraction_word_timeout_seconds=20, upload_max_bytes=10_000_000,
                  knowledge_max_pdf_pages=10, knowledge_max_extracted_characters=100_000,
                  knowledge_max_docx_uncompressed_bytes=10_000_000)
    values.update(overrides)
    return TestSettings(**values)


def pdf_bytes(pages: int = 1, text: str = "Value -12.50") -> bytes:
    with fitz.open() as document:
        for _ in range(pages):
            page = document.new_page(width=240, height=160)
            page.insert_text((20, 40), text)
        return document.tobytes()


def image_bytes(frames: int = 1) -> bytes:
    images = [Image.new("RGB", (120, 100), (250 - number, 250, 250)) for number in range(frames)]
    output = io.BytesIO()
    images[0].save(output, format="TIFF" if frames > 1 else "PNG", save_all=True, append_images=images[1:])
    return output.getvalue()


def provider(monkeypatch: pytest.MonkeyPatch, *, text: str = "Value -12.50", kind: str = "paragraph", truncate: bool = False) -> list[dict]:
    calls: list[dict] = []

    async def read(image: bytes, config: TestSettings, prompt: str) -> dict:
        calls.append({"image": hashlib.sha256(image).hexdigest(), "prompt": prompt})
        assert text not in prompt, "Native transcription leaked into independent vision prompt"
        raw = ({"regions": [{"kind": kind, "bbox": [0, 0, 1000, 1000]}], "complete": True}
               if prompt == regions.INVENTORY_PROMPT else {"kind": kind, "text": text, "complete": True})
        return {"model": MODEL, "raw": json.dumps(raw), "finish_reason": "length" if truncate else "stop", "usage": None}

    monkeypatch.setattr(regions, "read_evidence", read)
    return calls


@pytest.mark.asyncio
@pytest.mark.parametrize(("encoding", "bom"), [("utf-8", b""), ("utf-8", codecs.BOM_UTF8),
                          ("utf-16-le", codecs.BOM_UTF16_LE), ("utf-16-be", codecs.BOM_UTF16_BE),
                          ("utf-32-le", codecs.BOM_UTF32_LE), ("utf-32-be", codecs.BOM_UTF32_BE)])
async def test_strict_text_all_encodings_no_model(monkeypatch: pytest.MonkeyPatch, encoding: str, bom: bytes) -> None:
    calls = provider(monkeypatch)
    text = "Do NOT change −12.50.\r\nA\tB\n"  # Deliberately includes symbols and mixed line endings.
    original = bom + text.encode(encoding)
    package = await extract_document(original, "source.txt", settings())
    page = package.record["pages"][0]
    assert page["selected_text"] == text
    assert page["source_page_number"] is None
    assert page["lines"][-1]["byte_end"] == len(original)
    assert package.images["source/original.txt"] == original
    assert not calls
    assert not package.record["release_eligible"]


@pytest.mark.asyncio
async def test_bad_encoding_exports_a_blocked_draft() -> None:
    package = await extract_document(b"bad\xff", "source.txt", settings())
    assert package.issues[0]["code"] == "ENCODING_REQUIRED"
    assert "BLOCKED" in package.validation_text()
    assert package.archive()


@pytest.mark.asyncio
async def test_pdf_native_independent_visual_inputs_and_archive(monkeypatch: pytest.MonkeyPatch) -> None:
    calls = provider(monkeypatch)
    original = pdf_bytes()
    package = await extract_document(original, "source.pdf", settings())
    page = package.record["pages"][0]
    assert page["selected_text"] == "Value -12.50"
    assert len(calls) == 2
    assert page["source_inventory"][0]["disposition"] == "ALIGNED"
    assert page["source_page_index"] == 0
    assert page["native_words"]
    assert page["blocks"][0]["crop"] in package.images
    assert all(attempt["image_artifact"] in package.images for attempt in page["vision_attempts"])
    with zipfile.ZipFile(io.BytesIO(package.archive())) as archive:
        manifest = json.loads(archive.read("manifest.json"))
        for name, checksum in manifest["artifacts"].items():
            assert hashlib.sha256(archive.read(name)).hexdigest() == checksum
        assert archive.read("source/original.pdf") == original
    assert not package.record["release_eligible"]


@pytest.mark.asyncio
async def test_critical_sign_conflict_never_selects_a_winner(monkeypatch: pytest.MonkeyPatch) -> None:
    provider(monkeypatch, text="Value 12.50")
    package = await extract_document(pdf_bytes(), "source.pdf", settings())
    assert any(issue["code"] == "NUMERIC_CONFLICT" for issue in package.issues)
    assert package.record["pages"][0]["selected_text"] is None


@pytest.mark.asyncio
async def test_truncated_vision_is_bounded_and_cannot_checkpoint(monkeypatch: pytest.MonkeyPatch) -> None:
    calls = provider(monkeypatch, truncate=True)
    saved = []

    async def checkpoint(*values: object) -> None:
        saved.append(values)

    package = await extract_document(pdf_bytes(), "source.pdf", settings(), on_page=checkpoint)
    assert 1 <= len(calls) <= 16
    assert not saved
    assert not package.record["pages"][0]["inventory_complete"]
    assert any(issue.get("technical") for issue in package.issues)
    assert package.archive()


@pytest.mark.asyncio
async def test_completed_page_checkpoint_roundtrip_and_resume(monkeypatch: pytest.MonkeyPatch) -> None:
    calls = provider(monkeypatch)
    original = pdf_bytes(2)
    source_hash = hashlib.sha256(original).hexdigest()
    saved = {}

    async def checkpoint(page: dict, image: bytes, issues: list[dict]) -> None:
        blob = encode_page_checkpoint(page, image, issues, source_sha256=source_hash, model=MODEL, version=1)
        saved[page["number"]] = decode_page_checkpoint(blob, source_sha256=source_hash, model=MODEL, version=1, page_number=page["number"])
        with pytest.raises(ValueError):
            decode_page_checkpoint(blob, source_sha256="0" * 64, model=MODEL, version=1, page_number=page["number"])

    first = await extract_document(original, "source.pdf", settings(), on_page=checkpoint)
    assert len(saved) == 2
    calls.clear()
    resumed = await extract_document(original, "source.pdf", settings(), resume_pages=saved)
    assert not calls
    assert resumed.record["reused_units"] == [1, 2]
    assert first.record["pages"] == resumed.record["pages"]
    calls.clear()
    partial = await extract_document(original, "source.pdf", settings(), resume_pages={2: saved[2]})
    assert partial.record["reused_units"] == [2]
    assert partial.record["reprocessed_units"] == [1]
    assert len(calls) == 2


@pytest.mark.asyncio
async def test_corrupted_checkpoint_image_rerenders_native_source(monkeypatch: pytest.MonkeyPatch) -> None:
    calls = provider(monkeypatch)
    original = pdf_bytes()
    first = await extract_document(original, "source.pdf", settings())
    page = first.record["pages"][0]
    cache = {1: (deepcopy(page), b"not a valid image", deepcopy(first.issues))}
    calls.clear()
    second = await extract_document(original, "source.pdf", settings(), resume_pages=cache)
    assert len(calls) == 2
    assert second.record["pages"][0]["native_text"].strip() == "Value -12.50"
    assert second.record["reprocessed_units"] == [1]


@pytest.mark.asyncio
@pytest.mark.parametrize("suffix", ["pdf", "png"])
async def test_storage_failures_propagate_for_job_retry(monkeypatch: pytest.MonkeyPatch, suffix: str) -> None:
    provider(monkeypatch)

    async def checkpoint(*_: object) -> None:
        raise RuntimeError("simulated storage outage")

    with pytest.raises(CheckpointWriteError):
        await extract_document(pdf_bytes() if suffix == "pdf" else image_bytes(), f"source.{suffix}", settings(), on_page=checkpoint)


@pytest.mark.asyncio
async def test_multiframe_tiff_preserves_frame_identity(monkeypatch: pytest.MonkeyPatch) -> None:
    provider(monkeypatch)
    package = await extract_document(image_bytes(2), "source.tiff", settings())
    assert package.record["expected_units"] == 2
    assert [page["source_frame_index"] for page in package.record["pages"]] == [0, 1]
    assert all(page["source_page_number"] is None for page in package.record["pages"])
    assert package.record["pages"][0]["image_sha256"] != package.record["pages"][1]["image_sha256"]


@pytest.mark.asyncio
async def test_encrypted_pdf_exports_password_blocker(monkeypatch: pytest.MonkeyPatch) -> None:
    calls = provider(monkeypatch)
    with fitz.open(stream=pdf_bytes(), filetype="pdf") as document:
        encrypted = document.tobytes(encryption=fitz.PDF_ENCRYPT_AES_256, user_pw="test-only", owner_pw="owner-only")
    package = await extract_document(encrypted, "encrypted.pdf", settings())
    assert package.record["review_status"] == "needs_password"
    assert not calls
    assert package.archive()


@pytest.mark.asyncio
async def test_previous_package_reuses_only_matching_config(monkeypatch: pytest.MonkeyPatch) -> None:
    calls = provider(monkeypatch)
    original = pdf_bytes(2)
    first = await extract_document(original, "source.pdf", settings())
    calls.clear()
    second = await extract_document(original, "source.pdf", settings(), previous_package=first)
    assert not calls
    assert second.record["reused_units"] == [1, 2]
    third = await extract_document(original, "source.pdf", settings(extraction_render_dpi=96), previous_package=first)
    assert len(calls) == 4
    assert third.record["reprocessed_units"] == [1, 2]


def docx_bytes() -> bytes:
    document = Document()
    document.add_paragraph("Value -12.50")
    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "Group"
    table.cell(0, 1).text = "Value"
    table.cell(1, 0).text = "A"
    table.cell(1, 1).text = "-12.50"
    stream = io.BytesIO()
    document.save(stream)
    return stream.getvalue()


def test_word_native_parts_table_cells_and_source_xml() -> None:
    original = docx_bytes()
    record, derivative = inspect_word(original, settings())
    assert record["parts"]
    assert record["blocks"]
    table = next(block for block in record["blocks"] if block["type"] == "table")
    assert [cell["text"] for cell in table["cells"]] == ["Group", "Value", "A", "-12.50"]
    assert "<w:tbl" in table["xml"]
    assert not record["render_blockers"]
    assert derivative != original


def test_word_xml_rejects_dtd_entities() -> None:
    with pytest.raises(ValueError):
        xml(b'<!DOCTYPE x [<!ENTITY x "unsafe">]><x>&x;</x>')


@pytest.mark.skipif(shutil.which("soffice") is None, reason="Controlled Word renderer not installed")
def test_controlled_word_renderer_records_provenance() -> None:
    _, derivative = inspect_word(docx_bytes(), settings())
    rendered, metadata = render_word(derivative, settings())
    with fitz.open(stream=rendered, filetype="pdf") as document:
        assert len(document) >= 1
        assert "-12.50" in document[0].get_text()
    assert metadata["field_policy"] == "freeze_cached_values"
    assert metadata["fonts_sha256"]
    assert metadata["derivative_sha256"] == hashlib.sha256(derivative).hexdigest()


@pytest.mark.asyncio
async def test_missing_word_renderer_keeps_native_draft(monkeypatch: pytest.MonkeyPatch) -> None:
    calls = provider(monkeypatch)
    package = await extract_document(docx_bytes(), "source.docx", settings(extraction_word_renderer="no-such-renderer"))
    assert package.record["word"]["blocks"]
    assert "WORD_RENDER_REQUIRED" in {issue["code"] for issue in package.issues}
    assert not calls
    assert package.archive()


@pytest.mark.parametrize("name", ["../escape", "/absolute", "x/../escape", "x\\escape"])
def test_archive_paths_are_rejected(name: str) -> None:
    buffer = io.BytesIO()
    with zipfile.ZipFile(buffer, "w") as archive:
        archive.writestr(name, b"data")
    with pytest.raises(ValueError):
        read_zip(buffer.getvalue(), max_members=10, max_expanded_bytes=1000)


def test_archive_duplicate_members_and_expansion_are_rejected() -> None:
    buffer = io.BytesIO()
    with zipfile.ZipFile(buffer, "w") as archive:
        archive.writestr("large", b"x" * 1024)
    with pytest.raises(ValueError):
        read_zip(buffer.getvalue(), max_members=10, max_expanded_bytes=100)
    with pytest.warns(UserWarning):
        with zipfile.ZipFile(buffer, "a") as archive:
            archive.writestr("large", b"other")
    with pytest.raises(ValueError):
        read_zip(buffer.getvalue(), max_members=10, max_expanded_bytes=10000)
