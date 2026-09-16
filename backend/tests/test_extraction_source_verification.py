"""Real native parsing/exports; only paid/network vision calls are scripted."""
from __future__ import annotations

import asyncio
import codecs
import io
import json
import zipfile
from copy import deepcopy
from dataclasses import asdict, dataclass

import fitz
import pytest
from PIL import Image
from pydantic import ValidationError

from northstar_api.services import extraction_regions as regions
from northstar_api.services import hybrid_extraction as hybrid
from northstar_api.services.extraction_evidence import (
    crop_image, deterministic_zip, read_zip, review_projection, sha256, verify_projection,
)
from northstar_api.services.extraction_review import apply_review
from northstar_api.services.extraction_schema import Cell, Inventory, Reading, parse_response
from northstar_api.services.extraction_validation import critical_page, disagreement, review_digest, table_errors
from northstar_api.services.extraction_word import inspect_word, visible_text, xml


@dataclass
class Settings:
    upload_max_bytes: int = 8 * 1024**2
    extraction_max_evidence_bytes: int = 64 * 1024**2
    extraction_vision_model: str = "meta/llama-3.2-11b-vision-instruct"
    extraction_vision_provider: str = "nvidia"
    extraction_region_retries: int = 2
    extraction_max_calls_per_page: int = 24
    extraction_render_dpi: int = 100
    extraction_max_image_pixels: int = 1_000_000
    extraction_word_renderer: str = "libreoffice"
    extraction_word_timeout_seconds: int = 30
    knowledge_max_pdf_pages: int = 20
    knowledge_max_docx_uncompressed_bytes: int = 20 * 1024**2
    knowledge_max_extracted_characters: int = 500_000

    def model_dump(self):
        return asdict(self)


def pdf_bytes(text="Value -12.50", pages=1):
    with fitz.open() as pdf:
        for _ in range(pages):
            pdf.new_page().insert_text((72, 72), text)
        return pdf.tobytes()


def png_bytes(size=(200, 100)):
    stream = io.BytesIO()
    Image.new("RGB", size, "white").save(stream, format="PNG")
    return stream.getvalue()


def script(monkeypatch, text="Value -12.50"):
    calls = []
    async def read(image, prompt, settings):
        calls.append((image, prompt))
        assert text not in prompt  # independent pass never receives the transcript
        value = {"regions": [{"kind": "paragraph", "bbox": [0, 0, 1000, 1000]}], "complete": True} if "Independently inventory" in prompt else {
            "kind": "paragraph", "text": text, "complete": True}
        return {"raw": json.dumps(value), "finish_reason": "stop", "model": settings.extraction_vision_model, "usage": {}}
    monkeypatch.setattr(regions, "_read_evidence", read)
    return calls


def payload(record, text_override=None):
    result = []
    for page in record["pages"]:
        blocks = []
        for block in page.get("blocks", []):
            if block["status"] == "split":
                continue
            structure = deepcopy(block.get("reviewed_structure") or block["visual"])
            if text_override is not None:
                structure["text"] = text_override
            blocks.append({"id": block["id"], "selected_text": structure["text"], "structure": structure,
                           "related_block_ids": block.get("related_block_ids", [])})
        text = "\n".join(b["selected_text"] for b in blocks) if blocks else page.get("selected_text") or ""
        result.append({"number": page["number"], "selected_text": text,
                       "reason": "Compared with source evidence", "structure_confirmed": True,
                       "blank_confirmed": not text, "blocks": blocks})
    decisions = [{"id": issue["id"], "action": "source_verified", "reason": "Compared with source evidence"}
                 for issue in record["issues"] if not issue.get("technical")]
    return result, decisions


def approve(package, reviewer):
    pages, decisions = payload(package.record)
    apply_review(package.record, pages, decisions, reviewer)
    package.refresh()


@pytest.fixture(scope="module")
def base_package():
    with pytest.MonkeyPatch.context() as patch:
        script(patch)
        return asyncio.run(hybrid.extract_document(pdf_bytes(), "fixture.pdf", Settings()))


@pytest.fixture
def package(base_package):
    return deepcopy(base_package)


def test_real_pdf_native_render_and_source_inventory(package):
    page = package.record["pages"][0]
    assert page["native_text"].strip() == "Value -12.50"
    assert page["selected_text"] == "Value -12.50"
    assert page["native_words"] and page["source_page_number"] == 1
    assert page["vision_attempts"][0]["role"] == "VISION_PRIMARY"
    assert not package.record["release_eligible"]
    assert len(package.record["checks"]) == 10
    assert package.record["metrics"]["measured_transcription_accuracy"] == "NOT_MEASURED"


def test_numeric_release_needs_two_independent_current_reviews(package):
    approve(package, "reviewer-a")
    assert not package.record["release_eligible"]
    approve(package, "reviewer-a")
    assert not package.record["release_eligible"]
    approve(package, "reviewer-b")
    assert package.record["release_eligible"], package.record["checks"]
    package.record["review_status"] = "verified_release"
    assert "Verified release: APPROVED" in package.validation_text()


def test_text_without_critical_content_needs_one_source_review():
    package = asyncio.run(hybrid.extract_document(b"Hello\r\nworld\n", "text.txt", Settings()))
    assert all(s["status"] == "not_applicable" for s in package.record["stages"] if s["stage"] in {"render", "visual_inventory", "region_reading"})
    approve(package, "a")
    assert package.record["release_eligible"], package.record["checks"]


def test_atomic_review_keeps_record_unchanged_on_failure(package):
    before = deepcopy(package.record)
    pages, decisions = payload(package.record)
    pages[0]["blocks"][0]["structure"]["uncertain"] = True
    with pytest.raises(ValueError):
        apply_review(package.record, pages, decisions, "a")
    assert package.record == before


def test_api_captured_page_list_reference_remains_current(package):
    captured = package.record["pages"]
    approve(package, "a")
    assert captured is package.record["pages"]
    assert captured[0].get("primary_review")


@pytest.mark.parametrize("mutation", ["selected", "geometry", "links", "caption", "order", "candidate", "image", "type"])
def test_source_mutations_invalidate_approval(package, mutation):
    approve(package, "a"); approve(package, "b")
    page = package.record["pages"][0]
    before = review_digest(page)
    block = page["blocks"][0]
    if mutation == "selected": block["selected_text"] = "Value 12.50"
    if mutation == "geometry": block["bbox_px"][0] += 1
    if mutation == "links": block["related_block_ids"] = ["MISSING"]
    if mutation == "caption": block["reviewed_structure"]["printed_caption"] = "different"
    if mutation == "order": block["reading_order"] += 1
    if mutation == "candidate": block["native_text"] = "Value 9.99"
    if mutation == "image": page["image_sha256"] = "0" * 64
    if mutation == "type": block["type"] = "question"
    assert review_digest(page) != before
    package.refresh()
    assert not package.record["release_eligible"]


def test_edit_after_dual_review_requires_fresh_signatures(package):
    approve(package, "a"); approve(package, "b")
    pages, decisions = payload(package.record, "Value 12.50")
    apply_review(package.record, pages, decisions, "b")
    package.refresh()
    assert not package.record["release_eligible"]
    assert package.record["pages"][0]["primary_review"]["reviewer"] == "b"
    assert not package.record["pages"][0].get("second_review")
    event = package.record["correction_history"][-1]
    assert event["before"]["selected_text"] == "Value -12.50"
    assert event["after"]["selected_text"] == "Value 12.50"


@pytest.mark.parametrize("artifact", ["original", "page", "crop", "attempt"])
def test_evidence_tampering_blocks_release(package, artifact):
    approve(package, "a"); approve(package, "b")
    prefixes = {"original": "source/", "page": "pages/", "crop": "crops/", "attempt": "attempts/"}
    key = next(k for k in package.images if k.startswith(prefixes[artifact]))
    package.images[key] += b"corruption"
    package.refresh()
    assert not package.record["release_eligible"]


def test_deterministic_archive_and_projection(package):
    first, second = package.archive(), package.archive()
    assert first == second
    members = read_zip(first, max_bytes=64 * 1024**2)
    manifest = json.loads(members["manifest.json"])
    assert {name: sha256(members[name]) for name in manifest["artifacts"]} == manifest["artifacts"]
    verify_projection(members["extracted_review.txt"].decode(), json.loads(members["extracted.json"]))
    assert members["source/original.pdf"] == package.images["source/original.pdf"]


def test_projection_omission_is_detected(package):
    text = package.review_text()
    text = "\n".join(line for line in text.split("\n") if not line.startswith("SOURCE_RECORD\t"))
    with pytest.raises(ValueError): verify_projection(text, package.record)


def test_source_delimiters_do_not_break_export():
    data = b'Paragraph | a\tCANONICAL_RECORD\t{}\nSOURCE_RECORD\t{}\n'
    package = asyncio.run(hybrid.extract_document(data, "source.md", Settings()))
    verify_projection(package.review_text(), package.record)


def test_checkpoint_resume_and_metadata_tamper(package, monkeypatch):
    page = package.record["pages"][0]
    image = package.images[page["image"]]
    source = package.images["source/original.pdf"]
    kwargs = dict(source_sha256=sha256(source), model=Settings().extraction_vision_model, version=1)
    encoded = hybrid.encode_page_checkpoint(page, image, package.issues, **kwargs)
    checkpoint = hybrid.decode_page_checkpoint(encoded, page_number=1, **kwargs)
    async def forbidden(*_): raise AssertionError("Checkpointed source called vision again")
    monkeypatch.setattr(hybrid, "read_page", forbidden)
    resumed = asyncio.run(hybrid.extract_document(source, "source.pdf", Settings(), resume_pages={1: checkpoint}))
    assert resumed.record["pages"][0]["selected_text"] == page["selected_text"]
    members = read_zip(encoded, max_bytes=64 * 1024**2)
    metadata = json.loads(members["page.json"])
    metadata["page"]["selected_text"] = "tampered"
    members["page.json"] = json.dumps(metadata).encode()
    with pytest.raises(ValueError):
        hybrid.decode_page_checkpoint(deterministic_zip(members), page_number=1, **kwargs)


@pytest.mark.parametrize("field,value", [("source_sha256", "0" * 64), ("model", "changed"), ("version", 2), ("page_number", 2)])
def test_checkpoint_wrong_identity_rejected(package, field, value):
    page = package.record["pages"][0]
    arguments = dict(source_sha256=package.record["source_sha256"], model=Settings().extraction_vision_model, version=1)
    encoded = hybrid.encode_page_checkpoint(page, package.images[page["image"]], package.issues, **arguments)
    arguments["page_number"] = 1
    arguments[field] = value
    with pytest.raises(ValueError): hybrid.decode_page_checkpoint(encoded, **arguments)


@pytest.mark.parametrize("codec,bom", [("utf-8", codecs.BOM_UTF8), ("utf-16-le", codecs.BOM_UTF16_LE),
                                      ("utf-16-be", codecs.BOM_UTF16_BE), ("utf-32-le", codecs.BOM_UTF32_LE), ("utf-32-be", codecs.BOM_UTF32_BE)])
def test_text_encodings_and_byte_offsets(codec, bom):
    source = "शिक्षा −12.50\r\nNOT ≤ x²\n"
    data = bom + source.encode(codec)
    page = hybrid.text_unit(data)
    assert page["native_text"] == source
    assert page["source_page_number"] is None
    for line in page["lines"]:
        assert data[line["byte_start"]:line["byte_end"]].decode(codec) == line["text"]


def test_bad_encoding_keeps_blocked_draft():
    package = asyncio.run(hybrid.extract_document(b"bad\xff", "source.txt", Settings()))
    assert any(i["code"] == "ENCODING_REQUIRED" for i in package.issues)
    assert not package.record["release_eligible"]
    assert "extracted_review.txt" in read_zip(package.archive(), max_bytes=64 * 1024**2)


@pytest.mark.parametrize("raw", ['{"complete":true,"complete":false,"regions":[]}', '{"complete":true,"regions":[{"kind":"paragraph","bbox":[0,0,NaN,5]}]}'])
def test_duplicate_json_keys_and_nonfinite_values_rejected(raw):
    with pytest.raises(ValueError): parse_response(raw, Inventory)


@pytest.mark.parametrize("box", [[-1,0,10,10], [0,0,0,10], [0,0,201,50], [0,0,10.5,20]])
def test_crop_geometry_rejected(box):
    with pytest.raises(ValueError): crop_image(png_bytes(), box)


@pytest.mark.parametrize("name", ["../escape", "/absolute", "x/../y", "x\\y", "C:/file"])
def test_archive_paths_rejected(name):
    with pytest.raises(ValueError): deterministic_zip({name: b"text"})


def test_duplicate_zip_members_rejected():
    buffer = io.BytesIO()
    with zipfile.ZipFile(buffer, "w") as archive:
        archive.writestr("x", "a")
        with pytest.warns(UserWarning): archive.writestr("x", "b")
    with pytest.raises(ValueError): read_zip(buffer.getvalue(), max_bytes=100)


def test_table_grid_header_relationships_and_merge_anchor():
    cells = [dict(row=0,column=0,text="Year",header=True), dict(row=0,column=1,text="Value",header=True),
             dict(row=1,column=0,text="2026",header_refs=["R0C0"]), dict(row=1,column=1,text="-12.50",header_refs=["R0C1"])]
    assert table_errors(cells) == []
    cells[3]["header_refs"] = ["R0C0"]
    assert any("outside" in error for error in table_errors(cells))
    assert table_errors([dict(row=0,column=0,text="",state="COVERED_BY_MERGE")])
    assert not table_errors([dict(row=0,column=0,text="A",column_span=2), dict(row=0,column=1,text="",state="COVERED_BY_MERGE")])


@pytest.mark.parametrize("state", ["UNREADABLE", "AMBIGUOUS", "STRUCTURE_UNKNOWN"])
def test_unresolved_table_states_block(state):
    assert table_errors([dict(row=0,column=0,text="?",state=state)])


def test_table_resource_limit_and_duplicate_positions():
    with pytest.raises(ValidationError): Cell(row=0,column=0,text="x",row_span=1000,column_span=1000)
    assert table_errors([dict(row=0,column=0,text="A"), dict(row=0,column=0,text="B")])


@pytest.mark.parametrize("native,visual,code", [("NOT correct", "correct", "CRITICAL_WORD_CONFLICT"),
    ("-12.50", "12.50", "NUMERIC_CONFLICT"), ("x² ≤ 1", "x₂ < 1", "CRITICAL_SYMBOL_CONFLICT"),
    ("a paragraph", "a", "CANDIDATE_CONFLICT")])
def test_critical_token_mutations(native, visual, code):
    assert disagreement(native, visual) == code


def test_technical_failure_cannot_be_approved(package):
    package.issues[0]["technical"] = True
    before = deepcopy(package.record)
    pages, _ = payload(package.record)
    decisions = [{"id": package.issues[0]["id"], "action": "source_verified", "reason": "read source"}]
    with pytest.raises(ValueError): apply_review(package.record, pages, decisions, "a")
    assert package.record == before


def test_vision_output_truncation_has_bounded_distinct_recovery(monkeypatch):
    calls = []
    async def failed(image, prompt, settings):
        calls.append(sha256(image))
        return {"raw": "{}", "finish_reason": "length"}
    monkeypatch.setattr(regions, "_read_evidence", failed)
    settings = Settings(extraction_max_calls_per_page=5, extraction_region_retries=3)
    result = asyncio.run(regions.read_page(png_bytes(), 1, settings))
    assert len(calls) <= 5
    assert not result["inventory_complete"]
    assert any(p["technical"] for p in result["problems"])
    signatures = [(a["stage"], tuple(a["bbox_px"])) for a in result["vision_attempts"]]
    assert len(signatures) == len(set(signatures))


def test_missing_pdf_page_render_keeps_all_source_units(monkeypatch):
    script(monkeypatch)
    original = hybrid.pdf_page
    async def fail_middle(data, index, settings):
        if index == 1: raise ValueError("deliberate renderer failure")
        return await original(data, index, settings)
    monkeypatch.setattr(hybrid, "pdf_page", fail_middle)
    package = asyncio.run(hybrid.extract_document(pdf_bytes(pages=3), "three.pdf", Settings()))
    assert [p["number"] for p in package.record["pages"]] == [1,2,3]
    assert package.record["pages"][1]["failure_code"] == "PAGE_RENDER_FAILED"
    assert package.record["pages"][2]["selected_text"] == "Value -12.50"
    assert next(s for s in package.record["stages"] if s["stage"] == "render")["status"] == "partial"


def test_multiframe_image_inventory(monkeypatch):
    script(monkeypatch, "Visible text")
    data = io.BytesIO()
    Image.new("RGB", (100,100), "white").save(data, format="TIFF", save_all=True, append_images=[Image.new("RGB", (100,100), "black")])
    package = asyncio.run(hybrid.extract_document(data.getvalue(), "frames.tiff", Settings()))
    assert package.record["expected_units"] == 2
    assert [p["source_frame_index"] for p in package.record["pages"]] == [0,1]
    assert not package.record["release_eligible"]


def test_xml_doctype_and_hidden_false_and_equation_text():
    with pytest.raises(ValueError): xml(b'<!DOCTYPE root [<!ENTITY x "bad">]><root>&x;</root>')
    value = xml(b'<w:p xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main" xmlns:m="http://schemas.openxmlformats.org/officeDocument/2006/math"><w:r><w:rPr><w:vanish w:val="false"/></w:rPr><w:t>Shown</w:t></w:r><w:del><w:r><w:t>Deleted</w:t></w:r></w:del><m:oMath><m:r><m:t>x</m:t></m:r></m:oMath></w:p>')
    assert visible_text(value) == "Shownx"


def test_word_native_nested_table_and_part_checksums():
    from docx import Document
    doc = Document(); doc.add_paragraph("Heading")
    table = doc.add_table(rows=2, cols=2)
    table.cell(0,0).merge(table.cell(1,0)).text = "Merged"
    table.cell(0,1).text = "Value"
    table.cell(1,1).add_table(rows=1, cols=1).cell(0,0).text = "Nested"
    output = io.BytesIO(); doc.save(output)
    record, derivative = inspect_word(output.getvalue(), Settings())
    assert len([b for b in record["blocks"] if b["type"] == "table"]) == 2
    assert any(c.get("row_span") == 2 for b in record["blocks"] if b["type"] == "table" for c in b["cells"])
    assert record["original_sha256"] == sha256(output.getvalue())
    assert derivative and record["parts"]


def test_failed_unit_checkpoint_resumes_without_losing_later_pages(monkeypatch):
    script(monkeypatch)
    source = pdf_bytes(pages=3)
    original = hybrid.pdf_page
    saved = {}
    kwargs = dict(source_sha256=sha256(source), model=Settings().extraction_vision_model, version=1)
    async def checkpoint(page, image, issues):
        data = hybrid.encode_page_checkpoint(page, image, issues, **kwargs)
        saved[page['number']] = hybrid.decode_page_checkpoint(data, page_number=page['number'], **kwargs)
    async def fail_middle(data, index, settings):
        if index == 1: raise ValueError('injected failure')
        return await original(data, index, settings)
    monkeypatch.setattr(hybrid, 'pdf_page', fail_middle)
    asyncio.run(hybrid.extract_document(source, 'three.pdf', Settings(), on_page=checkpoint))
    assert sorted(saved) == [1, 2, 3]
    assert saved[2][1] == b'' and saved[2][0]['image'] is None
    indices = []
    async def recovered(data, index, settings):
        indices.append(index)
        return await original(data, index, settings)
    monkeypatch.setattr(hybrid, 'pdf_page', recovered)
    resumed = asyncio.run(hybrid.extract_document(source, 'three.pdf', Settings(), resume_pages=saved))
    assert indices == [1]
    assert all(p['selected_text'] == 'Value -12.50' for p in resumed.record['pages'])


def test_image_checkpoint_storage_failure_propagates(monkeypatch):
    script(monkeypatch, 'Visible text')
    async def unavailable(*args):
        raise OSError('object store is unavailable')
    with pytest.raises(hybrid.CheckpointPersistenceError):
        asyncio.run(hybrid.extract_document(png_bytes(), 'scan.png', Settings(), on_page=unavailable))


@pytest.mark.parametrize('mutation', ['duplicate', 'omit'])
def test_page_manifest_mutations_block_release(package, mutation):
    approve(package, 'a'); approve(package, 'b')
    if mutation == 'duplicate': package.record['pages'].append(deepcopy(package.record['pages'][0]))
    else: package.record['pages'].clear()
    package.refresh()
    assert not package.record['release_eligible']


def test_question_option_consistency_rejects_omission():
    with pytest.raises(ValidationError):
        Reading(kind='question', text='1. Select NOT true', complete=True,
                option_labels=['A', 'B', 'C'], options=[{'label':'A','text':'a'}, {'label':'B','text':'b'}])


def test_source_anomaly_is_preserved_not_rewritten(package):
    pages, decisions = payload(package.record)
    decisions[0]['action'] = 'source_anomaly'
    decisions[0]['reason'] = 'The original prints this value; preserved without correction'
    apply_review(package.record, pages, decisions, 'a')
    assert package.issues[0]['disposition'] == 'source_anomaly'
    assert package.record['pages'][0]['selected_text'] == 'Value -12.50'


def test_benchmark_requires_ground_truth_and_detects_source_mutation(package):
    from northstar_api.services.extraction_benchmark import evaluate
    gold = deepcopy(package.record)
    with pytest.raises(ValueError): evaluate(package.record, gold)
    gold.update(human_adjudicated=True, adjudicator='SYNTHETIC_TEST_FIXTURE')
    result = evaluate(package.record, gold)
    assert result['character_error_rate'] == 0
    package.record['pages'][0]['selected_text'] = 'Value 12.50'
    result = evaluate(package.record, gold)
    assert result['character_error_rate'] > 0
    assert result['counts']['critical_token_errors'] > 0
    assert result['table_cell_exact_match'] is None


def test_blank_inventory_cannot_be_used_to_invent_text(monkeypatch):
    async def blank(*args):
        return {'blocks':[], 'vision_attempts':[], 'blank':True, 'inventory_complete':True,
                'visual_text':'', 'problems':[], 'prompt_version':regions.PROMPT_VERSION}
    monkeypatch.setattr(hybrid, 'read_page', blank)
    package = asyncio.run(hybrid.extract_document(png_bytes(), 'blank.png', Settings()))
    pages, decisions = payload(package.record)
    pages[0]['selected_text'] = 'fabricated content'
    with pytest.raises(ValueError): apply_review(package.record, pages, decisions, 'a')


def test_vision_review_keeps_dispute_open_and_does_not_select_candidate(monkeypatch):
    source = pdf_bytes()
    async def read(image, prompt, settings):
        if 'Independently inventory' in prompt:
            value = {'regions':[{'kind':'paragraph','bbox':[0,0,1000,1000]}],'complete':True}
        else:
            value = {'kind':'paragraph','text':'Value 12.50','complete':True}
        return {'raw':json.dumps(value),'finish_reason':'stop','model':settings.extraction_vision_model}
    monkeypatch.setattr(regions, '_read_evidence', read)
    package = asyncio.run(hybrid.extract_document(source, 'dispute.pdf', Settings()))
    block = package.record['pages'][0]['blocks'][0]
    assert block['review_candidate']['text'] == 'Value 12.50'
    assert block['selected_text'] is None
    assert any(a['role'] == 'VISION_REVIEW' for a in package.record['pages'][0]['vision_attempts'])
    assert any(i['code'] == 'NUMERIC_CONFLICT' and i['status'] == 'open' for i in package.issues)


def test_word_ambiguous_mapping_requires_explicit_source_links(package):
    package.record['word'] = {'blocks':[{'id':'W-test','part':'word/document.xml','type':'paragraph',
        'channel':'visible_candidate','text':'Value -12.50','rendered_pages':[],'mapping_status':'ambiguous'}]}
    approve(package, 'a'); approve(package, 'b')
    assert not package.record['release_eligible']
    pages, decisions = payload(package.record)
    pages[0]['blocks'][0]['structure']['source_logical_ids'] = ['W-test']
    apply_review(package.record, pages, decisions, 'a')
    approve(package, 'b')
    assert package.record['word']['blocks'][0]['rendered_pages'] == [1]
    assert package.record['release_eligible'], package.record['checks']


def test_word_render_native_vision_review_end_to_end(monkeypatch):
    import shutil
    from docx import Document
    if not shutil.which('libreoffice'):
        pytest.skip('Controlled Word renderer is not installed')
    script(monkeypatch, 'Hello world')
    document = Document(); document.add_paragraph('Hello world')
    output = io.BytesIO(); document.save(output)
    package = asyncio.run(hybrid.extract_document(output.getvalue(), 'lesson.docx', Settings()))
    assert package.record['pages'][0]['source_unit'] == 'word_rendered_page', package.issues
    assert package.record['pages'][0]['source_page_number'] is None
    assert package.record['word']['renderer']['version']
    assert 'derivatives/rendered.pdf' in package.images
    approve(package, 'source-reviewer')
    assert package.record['release_eligible'], package.record['checks']
