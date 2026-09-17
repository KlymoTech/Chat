"""Source fidelity, evidence integrity, and release mutation tests (no provider calls)."""

import codecs
import hashlib
import io
import json
import zipfile
from copy import deepcopy
from unittest.mock import AsyncMock

import fitz
import pytest
from docx import Document
from PIL import Image

from northstar_api.config import Settings
from northstar_api.services import extraction_regions as regions
from northstar_api.services import hybrid_extraction as hybrid
from northstar_api.services.extraction_review import apply_review
from northstar_api.services.extraction_schema import Reading
from northstar_api.services.extraction_validation import table_errors
from northstar_api.services.extraction_word import inspect_word


@pytest.fixture
def settings():
    return Settings(_env_file=None, extraction_render_dpi=72, extraction_region_retries=3)


@pytest.fixture
def reader(monkeypatch):
    async def probe(_):
        return {"status": "PASSED", "model": "test-provider-transport"}, b"probe"

    monkeypatch.setattr(hybrid, "verify_image_capability", probe)

    async def read(image, settings, prompt):
        if prompt == regions.INVENTORY_PROMPT:
            value = {"regions": [{"kind": "paragraph", "bbox": [0, 0, 1000, 1000]}], "complete": True}
        else:
            value = {"kind": "paragraph", "text": "Value -12.50", "complete": True}
        return {"raw": json.dumps(value), "finish_reason": "stop"}

    mock = AsyncMock(side_effect=read)
    monkeypatch.setattr(regions, "read_evidence", mock)
    return mock


def pdf_bytes(text="Value -12.50", pages=1):
    with fitz.open() as pdf:
        for _ in range(pages):
            page = pdf.new_page(width=400, height=300)
            page.insert_text((30, 40), text)
        return pdf.tobytes()


def submissions(record, numbers=None):
    result = []
    for page in record["pages"]:
        if numbers and page["number"] not in numbers:
            continue
        blocks = []
        for block in page.get("blocks", []):
            if block.get("status") == "split":
                continue
            structure = deepcopy(block.get("reviewed_structure") or block["visual"])
            blocks.append(
                {
                    "id": block["id"],
                    "selected_text": structure["text"],
                    "structure": structure,
                    "related_block_ids": block.get("related_block_ids", []),
                }
            )
        result.append(
            {
                "number": page["number"],
                "selected_text": "\n".join(b["selected_text"] for b in blocks)
                if blocks
                else page["native_text"],
                "reason": "Compared every value to the source",
                "structure_confirmed": True,
                "blank_confirmed": True,
                "blocks": blocks,
            }
        )
    return result


def resolutions(record):
    return [
        {"id": i["id"], "action": "source_verified", "reason": "Verified against preserved source"}
        for i in record["issues"]
        if not i.get("technical") and i.get("status") != "resolved"
    ]


async def test_independent_reading_and_two_reviewers(reader, settings):
    package = await hybrid.extract_document(pdf_bytes(), "source.pdf", settings)
    assert package.record["pages"][0]["selected_text"] == "Value -12.50"
    assert all("Value -12.50" not in call.args[2] for call in reader.call_args_list)
    assert not package.record["release_eligible"]
    apply_review(package.record, submissions(package.record), resolutions(package.record), "alice")
    package.refresh()
    assert not package.record["release_eligible"]
    apply_review(package.record, submissions(package.record), [], "alice")
    package.refresh()
    assert not package.record["release_eligible"]
    apply_review(package.record, submissions(package.record), [], "bob")
    package.refresh()
    assert package.record["release_eligible"], package.record["checks"]
    package.record["pages"][0]["blocks"][0]["selected_text"] = "Value 12.50"
    package.refresh()
    assert not package.record["release_eligible"]


async def test_review_is_atomic(reader, settings):
    package = await hybrid.extract_document(pdf_bytes(), "source.pdf", settings)
    before = deepcopy(package.record)
    review = submissions(package.record)
    review[0]["selected_text"] = "changed outside region"
    with pytest.raises(ValueError, match="projection"):
        apply_review(package.record, review, resolutions(package.record), "alice")
    assert package.record == before


async def test_partial_review_never_releases_unreviewed_pages(reader, settings):
    package = await hybrid.extract_document(pdf_bytes(pages=2), "source.pdf", settings)
    for reviewer in ["alice", "bob"]:
        apply_review(package.record, submissions(package.record, [1]), resolutions(package.record), reviewer)
    package.refresh()
    assert not package.record["release_eligible"]
    assert not package.record["pages"][1].get("primary_review")
    for reviewer in ["alice", "bob"]:
        apply_review(package.record, submissions(package.record, [2]), [], reviewer)
    package.refresh()
    assert package.record["release_eligible"]


@pytest.mark.parametrize(
    "encoding,bom",
    [
        ("utf-8", codecs.BOM_UTF8),
        ("utf-16-le", codecs.BOM_UTF16_LE),
        ("utf-16-be", codecs.BOM_UTF16_BE),
        ("utf-32-le", codecs.BOM_UTF32_LE),
    ],
)
async def test_text_bytes_offsets_and_delimiters(settings, encoding, bom):
    text = "α −12.50\r\n\tNOT x² | x₂\nCANONICAL_RECORD\t{}\r"
    raw = bom + text.encode(encoding)
    package = await hybrid.extract_document(raw, "notes.txt", settings)
    page = package.record["pages"][0]
    assert page["native_text"] == text
    assert page["source_page_number"] is None
    for line in page["lines"]:
        assert raw[line["byte_start"] : line["byte_end"]].decode(encoding) == line["text"]
    archive = package.archive()
    assert archive == package.archive()
    with zipfile.ZipFile(io.BytesIO(archive)) as z:
        manifest = json.loads(z.read("manifest.json"))
        for name, sha in manifest["artifacts"].items():
            assert hashlib.sha256(z.read(name)).hexdigest() == sha
        assert text in z.read("extracted_review.txt").decode()
    changed = submissions(package.record)
    changed[0]["selected_text"] = "corrected text"
    with pytest.raises(ValueError, match="bytes are authoritative"):
        apply_review(package.record, changed, [], "alice")


async def test_bad_encoding_keeps_draft(settings):
    package = await hybrid.extract_document(b"bad\xff", "notes.txt", settings)
    assert package.issues[0]["code"] == "ENCODING_REQUIRED"
    assert "BLOCKED" in package.validation_text()
    assert package.archive()


@pytest.mark.parametrize(
    "mutation,expected",
    [
        ("Value 12.50", "NUMERIC_CONFLICT"),
        ("Which is correct?", "CRITICAL_WORD_CONFLICT"),
        ("Stem\nA. Yes", "CANDIDATE_CONFLICT"),
        ("First paragraph", "CANDIDATE_CONFLICT"),
    ],
)
async def test_content_mutations_block(reader, settings, mutation, expected):
    native = {"NUMERIC_CONFLICT": "Value -12.50", "CRITICAL_WORD_CONFLICT": "Which is NOT correct?"}.get(
        expected
    )
    if native is None:
        native = "Stem\nA. Yes\nB. No" if mutation.startswith("Stem") else "First paragraph\nSecond paragraph"

    async def read(image, settings, prompt):
        obj = (
            {"regions": [{"kind": "paragraph", "bbox": [0, 0, 1000, 1000]}], "complete": True}
            if prompt == regions.INVENTORY_PROMPT
            else {"kind": "paragraph", "text": mutation, "complete": True}
        )
        return {"raw": json.dumps(obj), "finish_reason": "stop"}

    reader.side_effect = read
    package = await hybrid.extract_document(pdf_bytes(native), "source.pdf", settings)
    assert expected in {i["code"] for i in package.issues}
    assert not package.record["release_eligible"]


async def test_bounded_truncation_keeps_raw_output(reader, settings):
    reader.side_effect = None
    reader.return_value = {"raw": '{"regions":[', "finish_reason": "length"}
    package = await hybrid.extract_document(pdf_bytes(), "source.pdf", settings)
    assert reader.await_count <= 1 + settings.extraction_region_retries
    assert any(i["code"] == "VISUAL_INVENTORY_FAILED" for i in package.issues)
    assert package.record["pages"][0]["vision_attempts"][0]["raw"] == '{"regions":['
    assert package.archive()


async def test_failed_capability_prevents_document_calls(reader, settings, monkeypatch):
    async def failed(_):
        return {"status": "FAILED"}, b"probe"

    monkeypatch.setattr(hybrid, "verify_image_capability", failed)
    package = await hybrid.extract_document(pdf_bytes(), "source.pdf", settings)
    reader.assert_not_awaited()
    assert any(i["code"] == "VISION_CAPABILITY_FAILED" for i in package.issues)


async def test_checkpoint_resume_preserves_completed_pages(reader, settings):
    raw = pdf_bytes()
    saved = {}

    async def checkpoint(page, image, issues):
        data = hybrid.encode_page_checkpoint(
            page, image, issues, source_sha256=hashlib.sha256(raw).hexdigest(), model="nvidia", version=1
        )
        saved[page["number"]] = hybrid.decode_page_checkpoint(
            data,
            source_sha256=hashlib.sha256(raw).hexdigest(),
            model="nvidia",
            version=1,
            page_number=page["number"],
        )

    await hybrid.extract_document(raw, "source.pdf", settings, on_page=checkpoint)
    reader.reset_mock()
    package = await hybrid.extract_document(raw, "source.pdf", settings, resume_pages=saved)
    reader.assert_not_awaited()
    assert package.record["pages"][0]["selected_text"] == "Value -12.50"
    assert "crops/P0001-R0001.png" in package.images
    assert package.archive()


@pytest.mark.parametrize(
    "mutation",
    ["duplicate_page", "missing_page", "missing_crop", "wrong_image", "wrong_original", "broken_link"],
)
async def test_evidence_mutations(reader, settings, mutation):
    package = await hybrid.extract_document(pdf_bytes(), "source.pdf", settings)
    for reviewer in ["alice", "bob"]:
        apply_review(package.record, submissions(package.record), resolutions(package.record), reviewer)
    package.refresh()
    assert package.record["release_eligible"]
    page = package.record["pages"][0]
    if mutation == "duplicate_page":
        package.record["pages"].append(deepcopy(page))
    elif mutation == "missing_page":
        package.record["pages"].clear()
    elif mutation == "missing_crop":
        package.images.pop(page["blocks"][0]["crop"])
    elif mutation == "wrong_image":
        package.images[page["image"]] = b"wrong"
    elif mutation == "wrong_original":
        package.images["source/original.pdf"] = b"wrong"
    elif mutation == "broken_link":
        page["blocks"][0]["related_block_ids"] = ["UNKNOWN"]
    package.refresh()
    assert not package.record["release_eligible"]


def test_merged_table_and_headers():
    cells = [
        dict(row=0, column=0, column_span=2, text="Voltage", header=True),
        dict(row=0, column=1, text="", state="COVERED_BY_MERGE"),
        dict(row=1, column=0, text="A", header_refs=["R0C0"]),
        dict(row=1, column=1, text="-12.50", header_refs=["R0C0"]),
    ]
    assert not table_errors(cells)
    changed = deepcopy(cells)
    changed[-1]["header_refs"] = ["R0C1"]
    assert any("unknown header" in e for e in table_errors(changed))
    assert table_errors(cells[:-1])
    assert table_errors(cells + [dict(row=1, column=1, text="extra")])
    assert table_errors([dict(row=0, column=0, text="", state="COVERED_BY_MERGE")])


async def test_multiple_image_frames(reader, settings):
    out = io.BytesIO()
    Image.new("RGB", (100, 100), "white").save(
        out, format="TIFF", save_all=True, append_images=[Image.new("RGB", (100, 100), "black")]
    )
    package = await hybrid.extract_document(out.getvalue(), "scan.tiff", settings)
    assert [p["source_frame_index"] for p in package.record["pages"]] == [0, 1]
    assert not package.record["release_eligible"]


def test_word_nested_tables_revisions_hidden_and_equations(settings):
    doc = Document()
    doc.add_paragraph("Visible")
    table = doc.add_table(rows=1, cols=1)
    table.cell(0, 0).text = "Outer"
    table.cell(0, 0).add_table(rows=1, cols=1).cell(0, 0).text = "Inner"
    output = io.BytesIO()
    doc.save(output)
    with zipfile.ZipFile(io.BytesIO(output.getvalue())) as z:
        parts = {n: z.read(n) for n in z.namelist()}
    value = (
        parts["word/document.xml"]
        .decode()
        .replace(
            "</w:body>",
            "<w:p><w:del><w:r><w:delText>Deleted</w:delText></w:r></w:del><w:r><w:rPr><w:vanish/></w:rPr><w:t>Hidden</w:t></w:r><m:oMath><m:r><m:t>x²</m:t></m:r></m:oMath></w:p></w:body>",
        )
    )
    parts["word/document.xml"] = value.encode()
    altered = io.BytesIO()
    with zipfile.ZipFile(altered, "w") as z:
        for n, b in parts.items():
            z.writestr(n, b)
    record, derivative = inspect_word(altered.getvalue(), settings)
    assert len([b for b in record["blocks"] if b["type"] == "table"]) == 2
    assert {"deleted", "hidden"} <= {c["channel"] for c in record["channels"]}
    assert any(b["type"] == "equation" and "x²" in b["xml"] for b in record["blocks"])
    assert all("Deleted" not in b["text"] and "Hidden" not in b["text"] for b in record["blocks"])
    assert derivative != altered.getvalue()


def test_incomplete_structured_reading_cannot_be_reviewed():
    assert Reading(kind="paragraph", text="partial", complete=False).complete is False


def test_benchmark_detects_shared_channel_error():
    from northstar_api.services.extraction_benchmark import evaluate

    reference = {
        "source_sha256": "h",
        "adjudicated_by": "human-reference",
        "pages": [{"number": 1, "text": "NOT -12.50", "blocks": [{"id": "r", "text": "NOT -12.50"}]}],
    }
    extracted = {
        "source_sha256": "h",
        "pages": [
            {
                "number": 1,
                "native_text": "12.50",
                "visual_text": "12.50",
                "selected_text": "12.50",
                "blocks": [{"id": "r", "selected_text": "12.50"}],
            }
        ],
    }
    report = evaluate(reference, extracted)
    assert report["counts"]["critical_token_sequences_exact"] == 0
    assert report["counts"]["aligned_character_errors"] > 0
    assert report["block_recall"] == 1
    with pytest.raises(ValueError):
        evaluate({**reference, "adjudicated_by": ""}, extracted)


async def test_critical_structures_and_links_require_review(reader, settings):
    package = await hybrid.extract_document(pdf_bytes(), "source.pdf", settings)
    page = package.record["pages"][0]
    block = page["blocks"][0]
    block["type"] = "caption"
    block["visual"] = Reading(kind="caption", text="Value -12.50", complete=True).model_dump()
    for reviewer in ["alice", "bob"]:
        apply_review(package.record, submissions(package.record), resolutions(package.record), reviewer)
    package.refresh()
    assert not package.record["release_eligible"]
    assert any("caption_of" in d for c in package.record["checks"] for d in c["details"])


async def test_source_anomaly_never_rewrites_text(settings):
    package = await hybrid.extract_document(b"Source says 2+2=5.", "source.txt", settings)
    for reviewer in ["alice", "bob"]:
        apply_review(package.record, submissions(package.record), [], reviewer)
    package.refresh()
    assert package.record["release_eligible"]
    assert package.record["pages"][0]["selected_text"] == "Source says 2+2=5."


async def test_word_without_renderer_exports_blocked_logical_evidence(reader, settings, monkeypatch):
    monkeypatch.setattr(hybrid, "render_word", lambda *args: (_ for _ in ()).throw(ValueError("unavailable")))
    doc = Document()
    doc.add_paragraph("Kept exactly")
    stream = io.BytesIO()
    doc.save(stream)
    package = await hybrid.extract_document(stream.getvalue(), "source.docx", settings)
    assert package.record["pages"][0]["source_unit"] == "logical_document"
    assert any(i["code"] == "WORD_RENDER_REQUIRED" for i in package.issues)
    assert package.record["word"]["blocks"]
    assert not package.record["release_eligible"]
    assert package.archive()


async def test_background_batches_resume_without_rereading(reader, settings, monkeypatch):
    from types import SimpleNamespace
    from uuid import uuid4

    from northstar_api.models import JobStatus, KnowledgeKind, KnowledgeSource, KnowledgeStatus
    from northstar_api.services import ingestion

    settings = settings.model_copy(
        update={"background_dispatch_enabled": True, "extraction_pages_per_task": 1}
    )
    monkeypatch.setattr(ingestion, "get_settings", lambda: settings)
    raw = pdf_bytes(pages=3)
    source = KnowledgeSource(
        id=uuid4(),
        tenant_id=uuid4(),
        workspace_id=uuid4(),
        agent_id=uuid4(),
        name="Friendly book title",
        object_key="knowledge/book.pdf",
        kind=KnowledgeKind.FILE,
        checksum=hashlib.sha256(raw).hexdigest(),
        revision=1,
    )
    checkpoints = {}
    archives = []

    async def download(*args, **kwargs):
        return raw, {"content_type": "application/pdf"}

    async def save_checkpoint(*args, **kwargs):
        checkpoints[kwargs["page_number"]] = kwargs["data"]
        return "checkpoint"

    async def save_archive(tenant, source_id, data, **kwargs):
        archives.append(data)
        return f"extractions/{tenant}/{source_id}/v0001.zip"

    store = SimpleNamespace(
        download=download,
        put_extraction_page=save_checkpoint,
        load_extraction_pages=AsyncMock(side_effect=lambda *a, **k: dict(checkpoints)),
        put_extraction=save_archive,
    )
    monkeypatch.setattr(ingestion, "object_store", store)
    job = SimpleNamespace(status=JobStatus.QUEUED, attempts=0)
    for expected in range(1, 4):
        session = SimpleNamespace(scalar=AsyncMock(side_effect=[source, job]), commit=AsyncMock())
        with pytest.raises(ingestion.ExtractionYield):
            await ingestion.ingest_source_async(session, source.id)
        assert len(checkpoints) == expected
        assert job.status == JobStatus.QUEUED
    session = SimpleNamespace(scalar=AsyncMock(side_effect=[source, job]), commit=AsyncMock())
    await ingestion.ingest_source_async(session, source.id)
    assert source.status == KnowledgeStatus.NEEDS_REVIEW
    assert len(archives) == 1 and reader.await_count == 6
    restored = hybrid.ExtractionPackage.from_archive(archives[0], source_sha256=source.checksum, revision=1)
    assert len(restored.record["pages"]) == 3
    assert not restored.record["release_eligible"]


@pytest.mark.parametrize('filename', ['broken.pdf', 'broken.docx', 'broken.png'])
async def test_malformed_sources_keep_reviewable_draft(settings, filename):
    raw = b'not a supported document body'
    package = await hybrid.extract_document(raw, filename, settings)
    assert package.record['review_status'] == 'failed'
    assert package.issues[0]['code'] == 'SOURCE_PARSE_FAILED'
    assert not package.record['release_eligible']
    with zipfile.ZipFile(io.BytesIO(package.archive())) as archive:
        assert archive.read('source/original.' + filename.rsplit('.', 1)[-1]) == raw


async def test_word_derivative_page_identity_is_separate(reader, settings, monkeypatch):
    doc = Document()
    doc.add_paragraph('Value -12.50')
    raw = io.BytesIO()
    doc.save(raw)
    monkeypatch.setattr(hybrid, 'render_word', lambda *args: (pdf_bytes(), {'version': 'fixture-renderer', 'fonts': 'fixture'}))
    package = await hybrid.extract_document(raw.getvalue(), 'source.docx', settings)
    page = package.record['pages'][0]
    assert page['source_page_number'] is None
    assert page['rendered_page_number'] == 1
    assert page['blocks'][0]['source_page_number'] is None
    assert 'derivatives/rendered.pdf' in package.images
    for reviewer in ['alice', 'bob']:
        apply_review(package.record, submissions(package.record), resolutions(package.record), reviewer)
    package.refresh()
    assert package.record['release_eligible'], package.record['checks']


async def test_encrypted_pdf_has_password_state_and_draft(settings):
    with fitz.open() as pdf:
        pdf.new_page()
        raw = pdf.tobytes(encryption=fitz.PDF_ENCRYPT_AES_256, owner_pw='fixture-owner', user_pw='fixture-password')
    package = await hybrid.extract_document(raw, 'protected.pdf', settings)
    assert package.record['review_status'] == 'needs_password'
    assert package.issues[0]['code'] == 'NEEDS_PASSWORD'
    assert package.archive()
